from pathlib import Path

from typing import List, Dict, Any

from openpyxl import load_workbook

from app.core.config import EXTRACTED_DIR

class XLSXExtractor:

    def __init__(self, xlsx_path: str):

        self.xlsx_path = Path(xlsx_path)

    def get_sheet_count(self) -> int:

        workbook = load_workbook(

            self.xlsx_path,

            read_only=True,

            data_only=True,

        )

        try:

            return len(workbook.sheetnames)

        finally:

            workbook.close()

    def get_record_count(self) -> int:

        workbook = load_workbook(

            self.xlsx_path,

            read_only=True,

            data_only=True,

        )

        count = 0

        try:

            for worksheet in workbook.worksheets:

                for row in worksheet.iter_rows(values_only=True):

                    values = [

                        value

                        for value in row

                        if value is not None

                        and str(value).strip()

                    ]

                    if not values:

                        continue

                    first = str(values[0]).strip().lower()

                    if first in {

                        "no",

                        "no.",

                        "number",

                        "serial",

                        "id",

                    }:

                        continue

                    count += 1

            return count

        finally:

            workbook.close()

    def extract_images(self, job_id: str) -> List[Dict[str, Any]]:

        output_dir = EXTRACTED_DIR / job_id / "xlsx"

        output_dir.mkdir(parents=True, exist_ok=True)

        workbook = load_workbook(

            self.xlsx_path,

            data_only=False,

        )

        results = []

        try:

            for worksheet in workbook.worksheets:
                images = getattr(
                    worksheet,
                    "_images",
                    [],
                )
                mark_header = self._find_mark_header(worksheet)
                for image_index, image in enumerate(
                    images,
                    start=1,
                ):

                    try:

                        image_data = image._data()

                    except Exception:

                        continue

                    anchor = image.anchor

                    row_number = None

                    column_number = None

                    cell_reference = None

                    if hasattr(anchor, "_from"):
                        cell = anchor._from
                        row_number = cell.row + 1
                        column_number = cell.col + 1

                        cell_reference = (

                            f"{self._column_letter(column_number)}"
                            f"{row_number}"
                        )

                    # The portfolio stores the trademark image in the
                    # bordered cell below the "Mark" column. Ignore any
                    # decorative/header images elsewhere in a sheet.
                    if mark_header is not None:
                        header_row, mark_column = mark_header
                        if (
                            column_number != mark_column
                            or row_number is None
                            or row_number <= header_row
                        ):
                            continue

                    extension = "png"

                    if hasattr(image, "path"):

                        suffix = Path(

                            str(image.path)

                        ).suffix.lower()

                        if suffix:

                            extension = suffix.lstrip(".")

                    filename = (

                        f"{worksheet.title}_"

                        f"{image_index:04d}."

                        f"{extension}"

                    )

                    output_path = (

                        output_dir / filename

                    )

                    output_path.write_bytes(

                        image_data

                    )

                    width = 0

                    height = 0

                    try:

                        from PIL import Image

                        with Image.open(output_path) as pil_image:

                            width, height = (

                                pil_image.size

                            )

                    except Exception:

                        pass

                    # -------------------------------------------------

                    row_data = {}

                    if row_number is not None:

                        row_data = self._get_row_data(

                            worksheet,

                            row_number,

                        )

                    results.append({

                        "id":

                            f"xlsx_"

                            f"{worksheet.title}_"

                            f"{image_index}",

                        "source_type":

                            "xlsx",

                        "source_file":

                            self.xlsx_path.name,

                        "page":

                            None,

                        "sheet":

                            worksheet.title,

                        "row":

                            row_number,

                        "cell":

                            cell_reference,

                        "image_path":

                            str(output_path),

                        "bbox":

                            None,

                        "width":

                            width,

                        "height":

                            height,

                        "confidence":

                            1.0,

                        "row_data":

                            row_data,

                        "client_owner":

                            self._find_value(

                                row_data,

                                [

                                    "owner",

                                    "applicant",

                                    "company",

                                    "client",

                                    "name",

                                ],

                            ),

                        "application_no":

                            self._find_value(

                                row_data,

                                [

                                    "application no.",

                                    "application no",

                                    "application number",

                                    "application",

                                    "app no",

                                ],

                            ),

                        "phone":

                            self._find_value(

                                row_data,

                                [

                                    "phone",

                                    "telephone",

                                    "tel",

                                    "mobile",

                                    "phone number",

                                ],

                            ),

                        "mark":

                            self._find_value(

                                row_data,

                                [

                                    "mark",

                                    "trade mark",

                                    "trademark",

                                    "brand",

                                ],

                            ),

                    })

        finally:

            workbook.close()

        return results

    @staticmethod
    def _find_mark_header(worksheet) -> tuple[int, int] | None:
        for row in worksheet.iter_rows():
            for cell in row:
                value = str(cell.value or "").strip().casefold()
                if value in {"mark", "trade mark", "trademark"}:
                    return cell.row, cell.column
        return None

    def _get_row_data(

        self,

        worksheet,

        row_number: int,

    ) -> Dict[str, Any]:

        headers = []

        for cell in worksheet[1]:

            value = cell.value

            if value is None:

                headers.append("")

            else:

                headers.append(

                    str(value).strip()

                )

        row_values = worksheet[row_number]

        data = {}

        for index, cell in enumerate(

            row_values

        ):

            if index >= len(headers):

                continue

            header = headers[index]

            if not header:

                continue

            data[header] = cell.value

        return data

    @staticmethod

    def _find_value(

        row_data: Dict[str, Any],

        possible_names: List[str],

    ):

        normalized = {}

        for key, value in row_data.items():

            key_normalized = (

                str(key)

                .strip()

                .lower()

                .replace("_", " ")

                .replace("-", " ")

            )

            normalized[key_normalized] = value

        for possible in possible_names:

            possible_normalized = (

                possible

                .strip()

                .lower()

                .replace("_", " ")

                .replace("-", " ")

            )

            for key, value in normalized.items():

                if possible_normalized == key:

                    if value is not None:

                        return str(value).strip()

        return None

    @staticmethod

    def _column_letter(number: int) -> str:

        result = ""

        while number:

            number, remainder = divmod(

                number - 1,

                26,

            )

            result = (

                chr(65 + remainder)

                + result

            )

        return result